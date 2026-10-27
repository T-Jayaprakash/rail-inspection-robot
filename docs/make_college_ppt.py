"""Fills the college review template (docs/college_template.pptx) with this project's content.

Keeps the template's background, logos, fonts, colours and layout - only the text,
tables and pictures are replaced.  Output: docs/Rail_Inspection_Robot_Phase1.pptx
"""
import copy
import os
import re

from lxml import etree
from pptx import Presentation
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SS = os.path.join(HERE, "screenshots")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

prs = Presentation(os.path.join(HERE, "college_template.pptx"))
S = list(prs.slides)


def shape(slide, name):
    return next(x for x in slide.shapes if x.name == name)


def _runs(text):
    """'plain **bold** plain' -> [(text, bold)]"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") else (p, False) for p in parts if p]


def set_paras(tf, paras, size=None, keep_bold=False):
    """Replace all paragraphs, reusing the first paragraph's bullet/spacing and the
    first run's font so the template look is kept."""
    body = tf._txBody
    ps = body.findall(A + "p")
    tpl_p = ps[0]
    r0 = tpl_p.find(A + "r")
    tpl_rpr = copy.deepcopy(r0.find(A + "rPr")) if r0 is not None and r0.find(A + "rPr") is not None \
        else etree.SubElement(etree.Element(A + "r"), A + "rPr")
    for p in ps:
        body.remove(p)
    for text in paras:
        p = copy.deepcopy(tpl_p)
        for child in list(p):
            if child.tag != A + "pPr":
                p.remove(child)
        for seg, bold in _runs(text):
            r = etree.SubElement(p, A + "r")
            rpr = copy.deepcopy(tpl_rpr)
            if not keep_bold:
                if bold:
                    rpr.set("b", "1")
                elif "b" in rpr.attrib:
                    del rpr.attrib["b"]
            if size:
                rpr.set("sz", str(int(size * 100)))
            r.append(rpr)
            t = etree.SubElement(r, A + "t")
            t.text = seg
        body.append(p)


def set_text(slide, name, paras, **kw):
    set_paras(shape(slide, name).text_frame, paras if isinstance(paras, list) else [paras], **kw)


def replace_picture(slide, name, img):
    """Swap the picture, fitted inside the original box (aspect kept)."""
    old = shape(slide, name)
    from PIL import Image
    w, h = Image.open(img).size
    bw, bh = old.width, old.height
    k = min(bw / w, bh / h)
    nw, nh = int(w * k), int(h * k)
    pic = slide.shapes.add_picture(img, old.left + (bw - nw) // 2, old.top + (bh - nh) // 2, nw, nh)
    old._element.getparent().remove(old._element)
    return pic


# ----------------------------------------------------------------- 1. title
s = S[0]
set_text(s, "TextBox 6", "AUTONOMOUS RAIL TRACK INSPECTION ROBOT", keep_bold=True)
tp = shape(s, "Title 1").text_frame.paragraphs[1].runs
for r in tp[:3]:
    r.text = ""
g = shape(s, "TextBox 4").text_frame
gp = g.paragraphs
gp[1].runs[0].text = "VINCY LEENA"
gp[3].runs[0].text = "Department of Information Technology,"
p = shape(s, "TextBox 10").text_frame.paragraphs
names = ["GOUTHAMAN A (8115U23IT032)", "JAYAPRAKASH T (8115U23IT041)", "JOHN ADAN J R (8115U23IT046)", ""]
for para, n in zip(p[1:5], names):
    para.runs[0].text = n
    for r in para.runs[1:]:
        r.text = ""

# --------------------------------------------------------------- 3. objective
set_text(S[2], "Content Placeholder 2", [
    "The primary objective of this project is to **automate railway track inspection** using a low-cost "
    "autonomous rail-mounted robot that replaces slow and error-prone manual patrolling.",
    "The robot aims to **detect rail-head cracks** in real time using computer vision and grade them as minor, "
    "major or critical based on their length across the rail head.",
    "It continuously **measures the track gauge** (1676 mm Broad Gauge) and reports any widening or "
    "tightening beyond the permitted tolerance.",
    "It detects **obstacles on the track** using a front camera and ultrasonic sensors, slows down early "
    "and performs an emergency stop to avoid collision.",
    "Every defect is **geo-tagged with chainage, GPS and a photo**, shown live on a control-room "
    "dashboard and exported as an inspection report.",
])

# ------------------------------------------------------------ 4. introduction
set_text(S[3], "Content Placeholder 2", [
    "Indian Railways operates one of the largest rail networks in the world, and **track defects** such as "
    "rail cracks and gauge deviation are a major cause of **derailments**.",
    "Track inspection in India is still largely **manual** – keymen walk or push a trolley along the line "
    "and check a few kilometres per shift by eye.",
    "Small cracks are **easily missed**, records are paper-based, and obstructions on the line are found only "
    "when a train arrives.",
    "This project introduces an **autonomous inspection robot** that runs on the rails like a mini train, "
    "uses **cameras and sensors** to find cracks, gauge faults and obstacles, and reports them **live** to "
    "the control room.",
])

# ------------------------------------------------------ 6-9. literature survey
LIT = [
    ("2015", "Railway fastener inspection by real-time machine vision",
     "Ç. Aytekin, Y. Rezaeitabar, S. Dogru, I. Ulusoy", "Image processing, pixel & histogram features",
     "Real-time detection of missing fasteners from on-board cameras."),
    ("2016", "Deep convolutional neural networks for detection of rail surface defects",
     "S. Faghih-Roohi, S. Hajizadeh, A. Núñez, R. Babuska, B. De Schutter", "Deep CNN",
     "Classified rail-surface defect images taken from track video."),
    ("2017", "Deep multitask learning for railway track inspection",
     "X. Gibert, V. M. Patel, R. Chellappa", "Multitask deep CNN",
     "Joint detection of fastener and sleeper defects in track images."),
    ("2017", "A hierarchical extractor-based visual rail surface inspection system",
     "J. Gan, Q. Li, J. Wang, H. Yu", "Coarse + fine extractor",
     "Learned rail background to isolate surface defects with low false alarms."),
    ("2017", "A new rail inspection method based on deep learning using laser cameras",
     "Y. Santur, M. Karakose, E. Akin", "3D laser camera + deep learning",
     "Detected fracture, scouring and wear on the rail surface."),
    ("2020", "Object detection at level crossing using deep learning",
     "M. A. B. Fayyaz, C. Johnson", "Deep learning, CCTV + radar",
     "Detected obstacles at level crossings using existing CCTV."),
    ("2021", "Two deep learning networks for rail surface defect inspection of limited samples with "
             "line-level label", "D. Zhang, K. Song, Q. Wang, Y. He, X. Wen, Y. Yan", "Deep segmentation networks",
     "Reduced the labelled data needed for defect detection."),
    ("2021", "A defect detection method for rail surface and fasteners based on deep convolutional neural "
             "network", "D. Zheng, L. Li, S. Zheng et al.", "YOLOv5, Mask R-CNN, ResNet",
     "Located rails/fasteners, segmented surface defects on a high-speed line."),
    ("2021", "A review of vision-based on-board obstacle detection and distance estimation in railways",
     "D. Ristić-Durrant, M. Franke, K. Michels", "Survey: CV & AI methods",
     "Identified need for on-board obstacle detection with distance estimation."),
    ("2021", "Detection of safe passage for trains at rail level crossings using deep learning",
     "T. Pamuła, W. Pamuła", "Deep learning video analysis",
     "Checked if the train's path is clear with 0.98 recall."),
]


def fill_cell(cell, text, size=14):
    set_paras(cell.text_frame, [text], size=size)


rows_per = [(S[5], 0, 3), (S[6], 3, 3), (S[7], 6, 3), (S[8], 9, 1)]
for slide, start, n in rows_per:
    tbl = next(x for x in slide.shapes if x.has_table).table
    for i in range(n):
        yr, title, auth, algo, work = LIT[start + i]
        row = tbl.rows[i + 1]
        for c, v in enumerate([f"{start + i + 1}.", yr, title, auth, algo, work]):
            fill_cell(row.cells[c], v)

# --------------------------------------------------------- 10. existing system
set_text(S[9], "Content Placeholder 2", [
    "Track inspection is mostly done by **manual patrolling** – keymen walk the section or use a **push "
    "trolley** and check the rails visually with a hand gauge.",
    "Manual inspection covers only **a few kilometres per shift** and depends on the attention and "
    "experience of the inspector, so **small cracks are often missed**.",
    "**Ultrasonic flaw-detection trolleys** and **track recording cars** are expensive, few in number and run "
    "only periodically.",
    "Existing low-cost prototypes use **simple IR / ultrasonic threshold sensors** that give no crack size, "
    "no gauge value and **no photo evidence**.",
    "Obstacles on the line are **not detected in advance**, and defect records are mostly **paper-based** "
    "without exact location.",
])

# --------------------------------------------------------- 11. proposed system
set_text(S[10], "Content Placeholder 2", [
    "The proposed system is an **autonomous rail-mounted robot** that drives along the track by itself with "
    "closed-loop motor control.",
    "A **down-looking camera** with OpenCV detects **rail-head cracks**, rejects rust / grease stains and "
    "**measures track gauge** in every frame.",
    "A **front camera** and **ultrasonic sensors** detect obstacles – the robot **slows down** at 15 m and "
    "makes an **emergency stop** at 2 m, then resumes when the track is clear.",
    "Each defect is tagged with **chainage (encoder), GPS, severity and photo**; critical defects raise an "
    "**alert** and are paint-marked.",
    "A **live web dashboard** shows cameras, sensor values, defect map and event log, with **manual "
    "override** and **CSV / PDF report** generation.",
])

# ------------------------------------------------------ 12. system architecture
replace_picture(S[11], "Picture 2", os.path.join(HERE, "block_diagram.png"))

# ---------------------------------------------------------------- 13. modules
MODULES = [
    ("Track Vision Module", [
        "The Track Vision Module processes frames from the **down-looking track camera** mounted under the "
        "robot with an LED ring light.",
        "It **locates both rails** in every frame using column-brightness profiles and computes the **track "
        "gauge** between the inner faces of the rails.",
        "A **morphological black-hat filter** highlights thin dark lines across the rail head, and a **shape "
        "test** separates real cracks from round rust or grease stains.",
        "Cracks are graded as **MINOR, MAJOR or CRITICAL** based on how far they run across the rail head.",
    ]),
    ("Obstacle Detection Module", [
        "The Obstacle Detection Module uses the **front camera** to find objects such as boulders, logs, "
        "animals or workers inside the **track corridor**.",
        "The distance to the object is estimated from where it touches the ground using the **camera "
        "geometry**, giving early warning up to about 25 m.",
        "**Ultrasonic sensors** give an accurate distance at short range, and both readings are combined "
        "(**sensor fusion**) for a reliable stop decision.",
    ]),
    ("Robot Control Module", [
        "The Robot Control Module runs a **state machine**: Idle, Inspecting, Slowing, Obstacle Stop, Manual "
        "and Completed.",
        "It sets the robot speed – **cruise** during inspection, **reduced** when an obstacle is seen far away "
        "and **zero** when it is within 2 m.",
        "After the obstacle is removed and the track stays **clear for 2 seconds**, the robot **resumes "
        "automatically**.",
        "The operator can stop the robot or drive it **manually**, while the safety interlock stays active.",
    ]),
    ("Motor & Sensor Firmware Module", [
        "The **ESP32** controller drives the two DC geared motors through a **BTS7960 driver** using a "
        "**PI speed controller** with encoder feedback.",
        "It reads the **wheel encoder** (distance), **ultrasonic sensors**, **IR rail sensors**, **MPU6050 "
        "IMU** (vibration and tilt) and battery voltage.",
        "It performs a **hardware emergency stop** for obstacles, the **E-stop button** and a **1-second "
        "watchdog**, so the robot stops even if the Raspberry Pi fails.",
        "Data is exchanged with the Raspberry Pi as **JSON over USB serial**.",
    ]),
    ("Defect Logging & Alert Module", [
        "The Defect Logging Module converts per-frame detections into **unique defects** – a crack must be "
        "seen in at least 2 frames and nearby detections are merged.",
        "Each defect is tagged with **chainage (km)**, **GPS coordinates**, rail side, severity and a **photo "
        "close-up**.",
        "**Critical** cracks and obstructions trigger an **alert to the control room** and fire the "
        "**paint marker** on the track.",
    ]),
    ("Dashboard & Report Module", [
        "A **Flask web dashboard** streams both cameras live and shows speed, chainage, gauge, obstacle "
        "range, vibration and battery.",
        "It displays a **track defect map**, **gauge profile**, **event log** and a **defect register** with "
        "photos.",
        "The operator can **start, stop or manually drive** the robot from any phone or laptop over Wi-Fi.",
        "At the end of the run an **inspection report** is generated in **CSV, JSON and PDF** formats.",
    ]),
]
set_text(S[12], "Content Placeholder 2", [m[0] for m in MODULES])
for i, (name, body) in enumerate(MODULES):
    slide = S[13 + i]
    set_text(slide, "Title 1", f"Module {i + 1} – {name}", keep_bold=True)
    set_text(slide, "Content Placeholder 2", body)

# ------------------------------------------------------------------ outputs
OUT = [
    (S[20], "Content Placeholder 6", "LIVE CONTROL-ROOM DASHBOARD", "04_web_dashboard.png"),
    (S[21], "Content Placeholder 6", "CRACK DETECTION OUTPUT", "02_critical_crack.png"),
    (S[22], "Content Placeholder 5", "GAUGE MEASUREMENT OUTPUT", "g_25.png"),
    (S[23], "Content Placeholder 5", "OBSTACLE DETECTION OUTPUT", "03_obstacle_stop.png"),
    (S[24], "Content Placeholder 6", "INSPECTION REPORT & VALIDATION", "05_report.png"),
]
for slide, ph, title, img in OUT:
    set_text(slide, "TextBox 2", title, keep_bold=True)
    replace_picture(slide, ph, os.path.join(SS, img))

# video slide: embed the demo video (poster frame shown until played)
s = S[25]
old = shape(s, "Content Placeholder 5")
L, T, W, H = old.left, old.top, old.width, old.height
old._element.getparent().remove(old._element)
vh = int(W * 9 / 16)
if vh > H:
    vh, W = H, int(H * 16 / 9)
s.shapes.add_movie(os.path.join(ROOT, "output", "demo_inspection.mp4"), L, T + (H - vh) // 2, W, vh,
                   poster_frame_image=os.path.join(SS, "video_poster.png"), mime_type="video/mp4")

# ---------------------------------------------------------------- conclusion
set_text(S[26], "Content Placeholder 2", [
    "The project demonstrates that a **low-cost autonomous robot** can inspect railway track in real time "
    "and reduce dependence on manual patrolling.",
    "In the validation run on an 80 m section, the system detected **7 out of 7 cracks with zero false "
    "alarms**, found **2 out of 2 gauge faults** with **less than 1 mm** measurement error, and **stopped "
    "safely for 2 out of 2 obstacles**.",
    "Every defect is recorded with **exact location and photo evidence** and reported live to the control "
    "room.",
    "The system makes track inspection **faster, safer and more reliable** at a fraction of the cost of "
    "track recording cars.",
])

# ------------------------------------------------------- future enhancement
set_text(S[27], "Content Placeholder 2", [
    "Integrating **deep learning models (YOLOv8)** trained on rail defect datasets can improve detection of "
    "cracks, missing clips, bolts and broken sleepers.",
    "Adding **ultrasonic or eddy-current probes** will allow detection of **internal rail flaws** that are not "
    "visible on the surface.",
    "A **laser line profiler** can measure **rail wear** and give a complete 3D track geometry profile.",
    "**Solar charging**, a **LoRa / 4G** long-range link and integration with railway maintenance "
    "software will enable long unattended inspection runs.",
])

# --------------------------------------------------------------- references
REFS = [
    "Ç. Aytekin, Y. Rezaeitabar, S. Dogru, and I. Ulusoy, “Railway fastener inspection by real-time machine "
    "vision,” *IEEE Trans. Syst., Man, Cybern.: Syst.*, vol. 45, no. 7, pp. 1101–1107, 2015.",
    "S. Faghih-Roohi, S. Hajizadeh, A. Núñez, R. Babuska, and B. De Schutter, “Deep convolutional neural "
    "networks for detection of rail surface defects,” in *Proc. IJCNN*, Vancouver, Canada, 2016.",
    "X. Gibert, V. M. Patel, and R. Chellappa, “Deep multitask learning for railway track inspection,” "
    "*IEEE Trans. Intell. Transp. Syst.*, vol. 18, no. 1, pp. 153–164, 2017.",
    "J. Gan, Q. Li, J. Wang, and H. Yu, “A hierarchical extractor-based visual rail surface inspection "
    "system,” *IEEE Sensors J.*, vol. 17, no. 23, pp. 7935–7944, 2017.",
    "Y. Santur, M. Karakose, and E. Akin, “A new rail inspection method based on deep learning using laser "
    "cameras,” in *Proc. Int. Artificial Intelligence and Data Processing Symp. (IDAP)*, 2017.",
    "M. A. B. Fayyaz and C. Johnson, “Object detection at level crossing using deep learning,” "
    "*Micromachines*, vol. 11, no. 12, p. 1055, 2020.",
    "D. Zhang, K. Song, Q. Wang, Y. He, X. Wen, and Y. Yan, “Two deep learning networks for rail surface "
    "defect inspection of limited samples with line-level label,” *IEEE Trans. Ind. Informat.*, vol. 17, "
    "pp. 6731–6741, 2021.",
    "D. Zheng, L. Li, S. Zheng, X. Chai, S. Zhao, Q. Tong, J. Wang, and L. Guo, “A defect detection method "
    "for rail surface and fasteners based on deep convolutional neural network,” 2021, "
    "doi: 10.1155/2021/2565500.",
    "D. Ristić-Durrant, M. Franke, and K. Michels, “A review of vision-based on-board obstacle detection and "
    "distance estimation in railways,” *Sensors*, vol. 21, no. 10, p. 3452, 2021.",
    "T. Pamuła and W. Pamuła, “Detection of safe passage for trains at rail level crossings using deep "
    "learning,” *Sensors*, vol. 21, no. 18, p. 6281, 2021.",
    "Ministry of Railways, Government of India, *Indian Railways Permanent Way Manual (IRPWM)*, 2020.",
    "R. C. Gonzalez and R. E. Woods, *Digital Image Processing*, 4th ed. New York, NY, USA: Pearson, 2018.",
    "OpenCV, “Morphological transformations,” OpenCV documentation. [Online]. Available: "
    "https://docs.opencv.org",
    "Espressif Systems, *ESP32 Technical Reference Manual*. [Online]. Available: https://www.espressif.com",
    "Raspberry Pi Ltd., *Raspberry Pi 5 documentation*. [Online]. Available: https://www.raspberrypi.com",
]


def ref_paras(lo, hi):
    out = []
    for i in range(lo, hi):
        # *italic* markers are dropped (template references use plain + italic journal; keep it simple)
        out.append(f"[{i + 1}] " + REFS[i].replace("*", ""))
    return out


set_text(S[28], "Content Placeholder 2", ref_paras(0, 5), size=16)
set_text(S[29], "Content Placeholder 2", ref_paras(5, 10), size=16)
set_text(S[30], "Content Placeholder 2", ref_paras(10, 15), size=16)

out = os.path.join(HERE, "Rail_Inspection_Robot_Phase1.pptx")
prs.save(out)
print("written", out)
