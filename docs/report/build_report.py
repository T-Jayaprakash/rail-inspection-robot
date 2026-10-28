"""Builds TrackBot_Project_Phase1_Report.docx (+ PDF) in the layout of the college Phase-I report.

Layout copied from the reference: US Letter, Times New Roman 13 pt, justified, 1.5 line spacing,
one blank line between paragraphs, bold centred chapter headings, bold left section headings,
bold centred figure captions, page numbers bottom-right starting at Chapter 1.

    python docs/report/build_report.py
"""
import copy
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
import content as C  # noqa: E402

FONT = "Times New Roman"
BODY_PT = 13
OUT_DOCX = os.path.join(HERE, "TrackBot_Project_Phase1_Report.docx")


# ------------------------------------------------------------------ helpers
def set_font(run, size=BODY_PT, bold=None, italic=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=14, before=0, spacing=1.5, left=None, first=None, keep=False):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    if left is not None:
        pf.left_indent = Inches(left)
    if first is not None:
        pf.first_line_indent = Inches(first)
    if keep:
        pf.keep_with_next = True
    return p


def add_rich(p, text, size=BODY_PT, bold_all=False):
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        b = part.startswith("**")
        r = p.add_run(part[2:-2] if b else part)
        set_font(r, size, bold=b or bold_all)
    return p


class Report:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(BODY_PT)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        sec = self.doc.sections[0]
        self._page(sec)
        self.headings = []        # (level, number, title) -> TOC
        self.figs = []            # (number, caption)
        self.first = True

    @staticmethod
    def _page(sec):
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(0.79)
        sec.top_margin, sec.bottom_margin = Inches(0.88), Inches(0.9)
        sec.footer_distance = Inches(0.45)

    # -------------------------------------------------------------- blocks
    def para(self, text, **kw):
        p = fmt(self.doc.add_paragraph(), **kw)
        return add_rich(p, text)

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        fmt(p, after=0, spacing=1.0)

    def centered(self, text, size=BODY_PT, after=6, bold=True):
        p = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=after, keep=True)
        set_font(p.add_run(text), size, bold=bold)
        return p

    def chapter(self, no, title):
        if not self.first:
            self.page_break()
        self.first = False
        self.centered(f"CHAPTER {no}", after=4)
        self.centered(title, after=14)
        self.headings.append((0, no, title))

    def heading(self, text):
        p = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.LEFT, after=12, before=4, keep=True)
        m = re.match(r"([\d.]+)\s+(.*)", text)
        if m:
            set_font(p.add_run(m.group(1) + " "), 12.7, bold=True)
            set_font(p.add_run(m.group(2)), BODY_PT, bold=True)
            lvl = m.group(1).count(".")
            self.headings.append((lvl, m.group(1), m.group(2)))
        else:
            set_font(p.add_run(text), BODY_PT, bold=True)

    def numbered(self, items, plain=False):
        for i, it in enumerate(items, 1):
            p = fmt(self.doc.add_paragraph(), left=0.5, first=-0.25, after=4 if plain else 6,
                    spacing=1.15 if plain else 1.5,
                    align=WD_ALIGN_PARAGRAPH.LEFT if plain else WD_ALIGN_PARAGRAPH.JUSTIFY)
            set_font(p.add_run(f"{i}.\t"), 12.7)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
            if plain:
                set_font(p.add_run(it), 12.7)
            else:
                bold, text = it
                set_font(p.add_run(bold), BODY_PT, bold=True)
                set_font(p.add_run(" – " + text), BODY_PT)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(14)

    def bullets(self, items):
        for it in items:
            p = fmt(self.doc.add_paragraph(), left=0.5, first=-0.25, after=2, spacing=1.15,
                    align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
            set_font(p.add_run("•\t"), 12.8, name="Symbol" if False else FONT)
            set_font(p.add_run(it), BODY_PT)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(14)

    def plain(self, lines):
        for ln in lines:
            p = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.LEFT, after=6, spacing=1.15)
            set_font(p.add_run(ln), BODY_PT)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(14)

    def figure(self, key):
        path, cap, width = C.FIG[key]
        p = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=6, spacing=1.0, keep=True)
        p.add_run().add_picture(os.path.join(HERE, path), width=Inches(width))
        c = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=14, spacing=1.15)
        set_font(c.add_run(f"Fig {key} {cap}"), BODY_PT, bold=True)
        self.figs.append((key, cap))

    def lit(self, year, authors, title, summary):
        for label, val, after in (("Year: ", year, 6), ("Authors: ", authors, 6), ("Title: ", title, 6)):
            p = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.LEFT, after=after, spacing=1.5,
                    keep=True)
            set_font(p.add_run(label), BODY_PT, bold=True)
            set_font(p.add_run(val), BODY_PT)
        p = fmt(self.doc.add_paragraph(), after=14)
        set_font(p.add_run("Summary: "), BODY_PT, bold=True)
        set_font(p.add_run(summary), BODY_PT)

    # --------------------------------------------------------- page numbers
    def new_numbered_section(self):
        """Start the body: new section, page numbers bottom-right starting at 1."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page(sec)
        sec.footer.is_linked_to_previous = False
        pg = sec._sectPr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sec._sectPr.append(pg)
        pg.set(qn("w:start"), "1")
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = fp.add_run()
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            r._element.append(el)
        set_font(r, 11, name="Calibri")
        self.first = True

    # ------------------------------------------------------ front matter
    def table(self, rows, widths, header_bold=True, center_cols=(), bold_cols=(), size=BODY_PT, height=0.36):
        t = self.doc.add_table(rows=len(rows), cols=len(widths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Inches(widths[j])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (j in center_cols or i == 0) else WD_ALIGN_PARAGRAPH.LEFT
                txt, b = (val if isinstance(val, tuple) else (val, False))
                set_font(p.add_run(txt), size if i else 13,
                         bold=(i == 0 and header_bold) or b or j in bold_cols)
            t.rows[i].height = Inches(height)
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
        tblW = t._tbl.tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            t._tbl.tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(layout)
        return t

    def toc(self, pages):
        self.centered("TABLE OF CONTENTS", size=15.9, after=12)
        rows = [[("CHAPTER NO", True), ("", True), ("TOPIC", True), ("PAGE NO", True)]]
        for t in ("ABSTRACT", "LIST OF FIGURES", "LIST OF ABBREVIATIONS"):
            rows.append(["", "", (t, True), ""])
        for lvl, no, title in self.headings_all:
            key = (lvl, no) if no else (lvl, title)
            pg = str(pages.get(key, ""))
            if lvl == 0:
                rows.append([(no, True) if no.isdigit() else "", "", (title.upper(), True), pg])
            else:
                rows.append(["", no, title, pg])
        t = self.table(rows, [1.2, 0.8, 4.0, 0.9], center_cols=(0, 1, 3), size=15.9, height=0.3)
        self._no_borders(t)

    def lof(self, pages):
        self.page_break()
        self.centered("LIST OF FIGURES", size=15.9, after=12)
        rows = [[("FIGURE NO", True), ("FIGURE NAME", True), ("PAGE NO", True)]]
        for key, cap in self.figs_all:
            rows.append([key, cap, str(pages.get(("fig", key), ""))])
        t = self.table(rows, [1.1, 4.4, 0.9], center_cols=(0, 2), size=14.1, height=0.42)
        self._no_borders(t)

    def loa(self):
        self.page_break()
        self.centered("LIST OF ABBREVIATIONS", size=15.9, after=12)
        rows = [[("S.NO", True), ("ABBREVIATION", True), ("MEANING", True)]]
        for i, (a, m) in enumerate(C.ABBR, 1):
            rows.append([str(i), (a, True), m])
        t = self.table(rows, [0.9, 1.9, 3.6], center_cols=(0, 1), size=BODY_PT, height=0.55)
        self._no_borders(t)

    @staticmethod
    def _no_borders(t):
        tbl = t._tbl
        tblPr = tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{e}")
            el.set(qn("w:val"), "nil")
            b.append(el)
        tblPr.append(b)

    # ------------------------------------------------------------ appendix
    def code_appendix(self):
        self.page_break()
        self.centered("APPENDIX A : CODING PYTHON", after=4)
        self.centered("PYTHON CODE", after=10)
        self.headings.append((0, "", "APPENDIX A: CODING"))
        for path, title in C.CODE_FILES:
            p = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.LEFT, after=6, before=6, keep=True)
            set_font(p.add_run("# " + title), BODY_PT, bold=True)
            with open(os.path.join(ROOT, path)) as f:
                lines = f.read().rstrip("\n").split("\n")
            for ln in lines:
                indent = len(ln) - len(ln.lstrip(" "))
                p = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.LEFT, after=0, spacing=1.5,
                        left=0.1 * (indent // 4) if indent else None)
                set_font(p.add_run(ln.strip() if ln.strip() else ""), BODY_PT)

    def screenshots(self):
        self.page_break()
        self.centered("APPENDIX B : SCREENSHOTS", after=14)
        self.headings.append((0, "", "APPENDIX B: SCREENSHOT"))
        for i, (path, cap) in enumerate(C.SCREENSHOTS):
            if i and i % 2 == 0:
                self.page_break()
            p = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=6, spacing=1.0, keep=True)
            p.add_run().add_picture(os.path.join(HERE, path), width=Inches(5.0))
            c = fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=24, spacing=1.15)
            set_font(c.add_run(cap), BODY_PT, bold=True)

    def references(self):
        self.page_break()
        self.centered("REFERENCES", after=14)
        self.headings.append((0, "", "REFERENCES"))
        for i, (a, it, b) in enumerate(C.REFERENCES, 1):
            p = fmt(self.doc.add_paragraph(), after=10)
            set_font(p.add_run(f"[{i}] " + a), BODY_PT)
            set_font(p.add_run(it), BODY_PT, italic=True)
            set_font(p.add_run(b), BODY_PT)


def build(pages):
    # pass 1 collects headings/figures; the front matter needs them first, so collect from content
    rep = Report()
    rep.headings_all, rep.figs_all = collect()
    rep.toc(pages)
    rep.lof(pages)
    rep.loa()
    rep.new_numbered_section()
    for blk in C.BODY:
        kind = blk[0]
        if kind == "chapter":
            rep.chapter(blk[1], blk[2])
        elif kind == "h":
            rep.heading(blk[1])
        elif kind == "p":
            rep.para(blk[1])
        elif kind == "num":
            rep.numbered(blk[1])
        elif kind == "numplain":
            rep.numbered(blk[1], plain=True)
        elif kind == "bul":
            rep.bullets(blk[1])
        elif kind == "plain":
            rep.plain(blk[1])
        elif kind == "fig":
            rep.figure(blk[1])
        elif kind == "lit":
            rep.lit(*blk[1:])
    rep.code_appendix()
    rep.screenshots()
    rep.references()
    rep.doc.core_properties.title = "TrackBot - Autonomous Railway Track Inspection System"
    rep.doc.core_properties.author = "Gouthaman A, Jayaprakash T, John Adan J R"
    rep.doc.save(OUT_DOCX)


def collect():
    heads, figs = [], []
    for blk in C.BODY:
        if blk[0] == "chapter":
            heads.append((0, blk[1], blk[2]))
        elif blk[0] == "h":
            m = re.match(r"([\d.]+)\s+(.*)", blk[1])
            if m:
                heads.append((m.group(1).count("."), m.group(1), m.group(2)))
        elif blk[0] == "fig":
            figs.append((blk[1], C.FIG[blk[1]][1]))
    heads += [(0, "", "APPENDIX A: CODING"), (0, "", "APPENDIX B: SCREENSHOT"), (0, "", "REFERENCES")]
    return heads, figs


def to_pdf():
    subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", HERE, OUT_DOCX],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return OUT_DOCX[:-5] + ".pdf"


def find_pages(pdf):
    """Locate headings/figures in the rendered PDF -> printed page numbers (body starts at 1)."""
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf], capture_output=True,
                                                       text=True).stdout).group(1))
    texts = [subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", pdf, "-"],
                            capture_output=True, text=True).stdout for i in range(1, n + 1)]
    body0 = next(i for i, t in enumerate(texts) if re.search(r"CHAPTER 1\s*\n", t))
    heads, figs = collect()
    pages = {}
    start = body0
    for lvl, no, title in heads:
        if lvl == 0 and no:
            pat = rf"CHAPTER {no}\s*\n"
        elif lvl == 0:
            pat = {"APPENDIX A: CODING": r"APPENDIX A : CODING", "APPENDIX B: SCREENSHOT": r"APPENDIX B : SCREENSHOTS",
                   "REFERENCES": r"^\s*REFERENCES\s*$"}[title]
        else:
            pat = rf"^\s*{re.escape(no)}\s+{re.escape(title[:25])}"
        for i in range(start, n):
            if re.search(pat, texts[i], re.M):
                pages[(lvl, no) if no else (lvl, title)] = i - body0 + 1
                start = i
                break
    for key, cap in figs:
        for i in range(body0, n):
            if re.search(rf"Fig {re.escape(key)}\s", texts[i]):
                pages[("fig", key)] = i - body0 + 1
                break
    return pages


if __name__ == "__main__":
    build({})
    pages = find_pages(to_pdf())
    build(pages)           # second pass fills TOC / List of Figures page numbers
    pdf = to_pdf()
    print("DOCX:", OUT_DOCX)
    print("PDF :", pdf)
    print("pages:", {k: v for k, v in pages.items() if k[0] in (0, "fig")})
