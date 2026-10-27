"""Builds docs/Rail_Inspection_Robot_Review.pptx (pip install python-pptx)."""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SS = os.path.join(HERE, "screenshots")

BG = RGBColor(0x0D, 0x11, 0x17)
PANEL = RGBColor(0x16, 0x1C, 0x25)
TXT = RGBColor(0xE6, 0xEC, 0xF3)
MUTED = RGBColor(0x8C, 0x98, 0xA8)
CYAN = RGBColor(0x3C, 0xBE, 0xE6)
GREEN = RGBColor(0x3F, 0xC8, 0x78)
RED = RGBColor(0xEB, 0x40, 0x40)
AMBER = RGBColor(0xF5, 0xAA, 0x28)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(title=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if kicker:
        text(s, kicker, 0.6, 0.35, 12, 0.4, 14, CYAN, bold=True)
    if title:
        text(s, title, 0.6, 0.65, 12, 0.8, 32, TXT, bold=True)
    return s


def text(s, t, x, y, w, h, size=18, color=TXT, bold=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
    return tb


def bullets(s, items, x, y, w, h, size=18):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        head, _, rest = it.partition(" - ")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = "▸ " + head
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), True, TXT
        if rest:
            r2 = p.add_run()
            r2.text = "  —  " + rest
            r2.font.size, r2.font.color.rgb = Pt(size - 3), MUTED
    return tb


def card(s, x, y, w, h, big, label, color=GREEN):
    sh = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = PANEL
    sh.line.color.rgb = RGBColor(0x2D, 0x37, 0x46)
    text(s, big, x + 0.2, y + 0.15, w - 0.4, 0.8, 34, color, bold=True)
    text(s, label, x + 0.2, y + 0.95, w - 0.4, 0.5, 13, MUTED)


def image(s, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(path, Inches(x), Inches(y), **kw)


GUIDE = "Vincy Leena"
TEAM = [("Gouthaman A", "8115U23IT032"), ("Jayaprakash T", "8115U23IT041"), ("John Adan J R", "8115U23IT046")]


def names_block(s, y):
    text(s, "Guided by :", 0.8, y, 4, 0.5, 18, CYAN, True)
    text(s, GUIDE, 0.8, y + 0.45, 5, 0.5, 20, TXT, True)
    text(s, "Presented by :", 7.4, y, 5, 0.5, 18, CYAN, True)
    for i, (n, r) in enumerate(TEAM):
        text(s, f"{n}  ({r})", 7.4, y + 0.45 + i * 0.45, 5.6, 0.5, 18, TXT, True)


# 1. title
s = slide()
bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.18))
bar.fill.solid(); bar.fill.fore_color.rgb = CYAN; bar.line.fill.background()
text(s, "FINAL YEAR PROJECT REVIEW", 0.8, 0.8, 12, 0.5, 16, CYAN, True)
text(s, "AUTONOMOUS RAIL TRACK", 0.8, 1.35, 12, 0.9, 44, TXT, True)
text(s, "INSPECTION ROBOT", 0.8, 2.15, 12, 0.9, 44, TXT, True)
text(s, "Crack detection  |  Gauge measurement  |  Obstacle detection  |  Live dashboard",
     0.8, 3.05, 12, 0.5, 18, MUTED)
line = s.shapes.add_shape(1, Inches(0.8), Inches(3.9), Inches(11.7), Inches(0.03))
line.fill.solid(); line.fill.fore_color.rgb = RGBColor(0x2D, 0x37, 0x46); line.line.fill.background()
names_block(s, 4.25)

# 2. abstract
s = slide("Abstract", "OVERVIEW")
text(s, ["Track defects such as rail cracks, gauge widening and obstructions are a major cause of derailments. "
         "In India, track inspection is still mostly manual: trackmen walk or push a trolley and check the "
         "rails by eye, which is slow and error-prone.",
         "",
         "This project presents a low-cost autonomous rail-mounted robot. It drives along the track by itself; a "
         "down-looking camera with OpenCV detects rail-head cracks and measures track gauge, while a front camera "
         "and ultrasonic sensors detect obstacles and stop the robot safely. Every defect is tagged with "
         "chainage (km), GPS, severity and a photo, shown live on a web dashboard, and exported as a report.",
         "",
         "On an 80 m validation run the system detected 7/7 cracks with 0 false alarms, 2/2 gauge faults "
         "(< 1 mm error) and stopped safely for 2/2 obstacles."], 0.8, 1.7, 11.8, 5.5, 18, TXT)

# 2. problem
s = slide("Problem: manual track inspection", "PROBLEM STATEMENT")
bullets(s, [
    "Trackmen walk / push a trolley - only a few km per shift, depends on human attention",
    "Small rail cracks are easily missed - cracks grow under load and lead to rail fracture",
    "Gauge widening is hard to judge by eye - needs a gauge measuring tool at every spot",
    "Obstructions on the line - boulders, fallen trees, cattle are found only when a train arrives",
    "Paper records - no exact location or photo evidence for the maintenance gang",
    "Track-recording cars are costly - few of them, so inspection is infrequent",
], 0.6, 1.7, 12, 5)

# existing system
s = slide("Existing system & its drawbacks", "EXISTING SYSTEM")
text(s, "Existing methods", 0.6, 1.6, 6, 0.5, 20, CYAN, True)
bullets(s, ["Manual patrolling - keyman walks the section daily",
            "Push trolley inspection - visual check, hand gauge",
            "Ultrasonic flaw-detection trolleys - operated manually",
            "Track recording cars - periodic, high cost"], 0.6, 2.2, 6, 4.5, 16)
text(s, "Drawbacks", 7.0, 1.6, 6, 0.5, 20, RED, True)
bullets(s, ["Slow - few km per shift", "Human error and fatigue", "No exact location / photo record",
            "Obstacles not detected in advance", "Infrequent, expensive equipment"], 7.0, 2.2, 5.8, 4.5, 16)

# 3. solution
s = slide("Proposed solution: an autonomous rail-mounted robot", "OUR IDEA")
bullets(s, [
    "Runs on the rails by itself - flanged wheels, DC motors, PI speed control",
    "Down-looking camera - detects rail-head cracks and measures track gauge in every frame",
    "Front camera + ultrasonic - slows down for far obstacles, emergency-stops at 2 m",
    "Encoder + GPS - every defect tagged with chainage (km) and coordinates",
    "Live web dashboard - control room sees cameras, defect map, alerts; can override",
    "Automatic report - CSV / PDF with photo of every defect, paint mark on critical ones",
], 0.6, 1.7, 7.2, 5, 17)
image(s, os.path.join(SS, "02_critical_crack.png"), 8.0, 1.9, w=4.9)

# 4. block diagram
s = slide("System block diagram", "HARDWARE ARCHITECTURE")
image(s, os.path.join(HERE, "block_diagram.png"), 0.9, 1.5, w=11.5)

# 5. components
s = slide("Hardware components (≈ ₹ 22-25k)", "BILL OF MATERIALS")
rows = [("Component", "Purpose", "₹ approx"),
        ("Raspberry Pi 5 (4 GB) + SD", "Vision, control, dashboard", "7,500"),
        ("ESP32 DevKit", "Motors, sensors, safety stop", "450"),
        ("2x camera + LED ring light", "Track cam + front cam", "3,900"),
        ("HC-SR04 x2 / TF-Luna LiDAR", "Obstacle range", "200 / 2,000"),
        ("MPU6050, TCRT5000 x2, encoder", "Vibration, rail presence, distance", "900"),
        ("NEO-6M GPS, SIM800L GSM", "Location, SMS alert", "900"),
        ("2x 12 V geared motor + BTS7960", "Drive", "2,050"),
        ("Chassis + flanged wheels", "Mechanical", "3,000"),
        ("3S Li-ion 10 Ah + BMS + buck", "Power", "3,000"),
        ("E-stop, buzzer, beacon, paint relay", "Safety & marking", "700")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2)).table
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = tbl.cell(i, j)
        c.text = v
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(0x22, 0x2B, 0x38) if i == 0 else PANEL
        p = c.text_frame.paragraphs[0]
        p.font.size = Pt(15)
        p.font.bold = i == 0
        p.font.color.rgb = CYAN if i == 0 else TXT
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(4.8), Inches(4.8), Inches(2.1)

# modules
s = slide("Project modules", "MODULE DESCRIPTION")
bullets(s, [
    "1. Track vision module - rail localisation, gauge measurement, crack detection (OpenCV)",
    "2. Obstacle detection module - front camera + ultrasonic sensor fusion",
    "3. Robot control module - state machine: inspect / slow / stop / resume / manual",
    "4. Motor & sensor firmware - ESP32: PI speed control, encoder, IMU, hard-stop, watchdog",
    "5. Defect logging module - de-duplication, chainage + GPS tagging, photo capture, alerts",
    "6. Dashboard & report module - Flask live web UI, defect map, CSV / PDF report",
], 0.6, 1.7, 12.2, 5.4, 18)

# software requirements
s = slide("Software requirements", "SOFTWARE")
rows = [("Item", "Used"),
        ("Programming languages", "Python 3, C++ (Arduino), JavaScript, HTML/CSS"),
        ("Computer vision", "OpenCV, NumPy"),
        ("Web framework", "Flask (MJPEG streaming, REST API)"),
        ("Firmware IDE", "Arduino IDE with ESP32 core 3.x"),
        ("Robot OS", "Raspberry Pi OS (64-bit)"),
        ("Communication", "USB serial (JSON), Wi-Fi / 4G"),
        ("Reports & demo", "Pillow, FFmpeg, CSV / HTML-PDF")]
tbl = s.shapes.add_table(len(rows), 2, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.4)).table
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = tbl.cell(i, j)
        c.text = v
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(0x22, 0x2B, 0x38) if i == 0 else PANEL
        p = c.text_frame.paragraphs[0]
        p.font.size = Pt(17)
        p.font.bold = i == 0 or j == 0
        p.font.color.rgb = CYAN if i == 0 else TXT
tbl.columns[0].width, tbl.columns[1].width = Inches(4.2), Inches(7.5)

# 6. algorithm
s = slide("Crack & gauge detection algorithm", "COMPUTER VISION (OpenCV)")
bullets(s, [
    "1. Rail localisation - 8 image bands, column-brightness high-pass → bright 72 mm runs",
    "2. Gauge - (centre R − centre L) × mm/px − head width, median of bands; ±10 mm alarm",
    "3. Black-hat filter - closing(img) − img lights up thin dark lines across the rail",
    "4. Shape test - length ≥ 35 % of head, aspect ≥ 1.8 → rejects round rust / grease stains",
    "5. Severity - MINOR ≥ 35 %, MAJOR ≥ 55 %, CRITICAL ≥ 80 % of rail head",
    "6. Tracking - seen in ≥ 2 frames, merged within 30 cm, located by encoder + pixel offset",
    "7. Obstacles - colour anomaly in track corridor, distance d = f·h / (y − y_horizon)",
], 0.6, 1.7, 12.2, 5.4, 17)

# 7. flowchart
s = slide("Inspection flowchart", "SOFTWARE LOGIC")
image(s, os.path.join(HERE, "flowchart.png"), 4.2, 1.35, h=6.0)
text(s, ["State machine:", "", "IDLE", "→ INSPECTING", "⇄ SLOWING (obstacle < 15 m)", "→ OBSTACLE_STOP (< 2 m)",
         "→ clear 2 s → INSPECTING", "→ COMPLETED", "", "Manual override + E-stop", "at any time"],
     0.6, 1.8, 3.4, 5, 16, MUTED)

# 8. dashboard
s = slide("Live control-room dashboard", "WEB APP (Flask)")
image(s, os.path.join(SS, "04_web_dashboard.png"), 0.6, 1.5, h=5.8)
text(s, ["Live track + front camera", "Speed, chainage, gauge,", "range, vibration, battery",
         "Defect map + gauge profile", "Event log + alerts", "Defect register with photos",
         "START / STOP / manual drive", "CSV + PDF report",
         "", "Opens on phone / laptop", "over Wi-Fi"], 7.0, 1.7, 5.8, 5.5, 17, TXT)

# 9. demo
s = slide("Demo: obstacle emergency stop", "RESULTS")
image(s, os.path.join(SS, "03_obstacle_stop.png"), 0.6, 1.5, w=12.1)

# 10. validation
s = slide("Validation results (80 m test section)", "RESULTS")
card(s, 0.6, 1.8, 2.3, 1.5, "7 / 7", "cracks detected")
card(s, 3.1, 1.8, 2.3, 1.5, "0", "false alarms (stains rejected)")
card(s, 5.6, 1.8, 2.3, 1.5, "2 / 2", "gauge faults")
card(s, 8.1, 1.8, 2.3, 1.5, "< 1 mm", "gauge RMS error")
card(s, 10.6, 1.8, 2.2, 1.5, "2 / 2", "obstacles stopped")
image(s, os.path.join(SS, "05_report.png"), 2.4, 3.6, h=3.7)

# 11. advantages / future
s = slide("Advantages, limitations & future scope", "DISCUSSION")
text(s, "Advantages", 0.6, 1.6, 4, 0.5, 20, GREEN, True)
bullets(s, ["Low cost", "Never tired, every metre recorded", "Exact location + photo",
            "Real-time alerts", "Fail-safe stop (ESP32 watchdog)"], 0.6, 2.2, 4, 4.5, 16)
text(s, "Limitations", 4.7, 1.6, 4, 0.5, 20, AMBER, True)
bullets(s, ["Surface cracks only (camera)", "Needs a maintenance block", "Rain / shadows → needs lighting"],
        4.7, 2.2, 4, 4.5, 16)
text(s, "Future scope", 8.8, 1.6, 4, 0.5, 20, CYAN, True)
bullets(s, ["YOLOv8 deep-learning detection", "Ultrasonic probe for internal flaws",
            "Laser rail-wear profile", "Missing clip / bolt detection", "Solar charging, LoRa link"],
        8.8, 2.2, 4.2, 4.5, 16)

# conclusion
s = slide("Conclusion", "CONCLUSION")
text(s, ["The project shows that a low-cost autonomous robot can inspect railway track in real time: it detects "
         "rail-head cracks with severity grading, measures track gauge to within 1 mm, and stops safely for "
         "obstacles.",
         "",
         "Every defect is recorded with exact chainage, GPS and photo evidence, and reported live to the "
         "control room. This makes track inspection faster, safer and more reliable than manual patrolling, "
         "at a fraction of the cost of track-recording cars."], 0.8, 1.8, 11.8, 5, 20, TXT)

# references
s = slide("References", "REFERENCES")
text(s, ["[1] Indian Railways Permanent Way Manual (IRPWM), Ministry of Railways, Government of India.",
         "[2] J. Gan, Q. Li, J. Wang, H. Yu, \"A Hierarchical Extractor-Based Visual Rail Surface Inspection "
         "System,\" IEEE Sensors Journal, vol. 17, no. 23, 2017.",
         "[3] R. C. Gonzalez, R. E. Woods, Digital Image Processing, Pearson (morphological processing).",
         "[4] OpenCV documentation - Morphological Transformations, https://docs.opencv.org",
         "[5] Ultralytics YOLOv8 documentation, https://docs.ultralytics.com",
         "[6] Espressif Systems, ESP32 Technical Reference Manual.",
         "[7] Raspberry Pi Foundation, Raspberry Pi 5 documentation."], 0.8, 1.7, 11.8, 5.5, 16, TXT)

# thank you
s = slide()
text(s, "Thank you", 0.8, 1.4, 12, 1.2, 54, TXT, True)
text(s, "Autonomous Rail Track Inspection Robot  —  Questions?", 0.8, 2.6, 12, 0.6, 22, MUTED)
names_block(s, 4.0)

out = os.path.join(HERE, "Rail_Inspection_Robot_Review.pptx")
prs.save(out)
print("written", out)
